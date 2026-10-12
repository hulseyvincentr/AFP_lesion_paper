#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
make_Q60_Q75_Q90_sensitivity_table.py

Sensitivity analysis for the AFP lesion phrase-duration result.

For every bird, calculate Q60, Q75, and Q90 of syllable-level
ΔCV = post-lesion CV - late-pre-lesion CV across all Figure-3-qualifying
syllables. Then compare bird-level quantile values among sham,
lateral-only, and pooled medial+lateral groups.

Primary-style tests:
  1) Medial+lateral > Sham          one-sided
  2) Medial+lateral > Lateral-only  one-sided
  3) Lateral-only vs Sham           two-sided

Permutation statistic: difference in MEAN bird-level quantile ΔCV.
Small sample sizes are handled by exact enumeration of bird-label assignments
when feasible.

Input must be supplied with --input (figure3_balanced_pair_metrics.csv).
Default output: outputs/Q60_Q75_Q90_sensitivity/

The script also writes a manuscript-ready Word document:
  05_Supplemental_Table_Q60_Q75_Q90_sensitivity.docx
"""

from __future__ import annotations

import argparse
import itertools
import math
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as exc:
    raise ImportError(
        "This script now writes a Word document and requires python-docx. "
        "Install it in the syntax_analysis environment with: conda install -c conda-forge python-docx"
    ) from exc

DEFAULT_OUTPUT = Path("outputs") / "Q60_Q75_Q90_sensitivity"


QUANTILES = (0.60, 0.75, 0.90)
SHAM = "sham_saline"
LATERAL = "lateral_only"
ML = "medial_and_lateral"
GROUP_ORDER = [SHAM, LATERAL, ML]
GROUP_LABEL = {SHAM: "Sham", LATERAL: "Lateral-only", ML: "Medial+lateral"}
COMPARISONS = [
    ("M+L vs Sham", ML, SHAM, "greater"),
    ("M+L vs Lateral-only", ML, LATERAL, "greater"),
    ("Lateral-only vs Sham", LATERAL, SHAM, "two-sided"),
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Q60/Q75/Q90 sensitivity analysis for bird-level phrase-duration ΔCV.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    p.add_argument("--input", type=Path, required=True)
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    p.add_argument(
        "--quantile-method", default="linear",
        choices=["linear", "lower", "higher", "midpoint", "nearest"]
    )
    p.add_argument("--max-exact-assignments", type=int, default=200_000)
    p.add_argument("--monte-carlo-permutations", type=int, default=100_000)
    p.add_argument("--seed", type=int, default=123)
    return p.parse_args()


def norm_text(x: object) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return " ".join(str(x).strip().lower().replace("_", " ").replace("-", " ").split())


def normalize_group(x: object) -> str:
    s = norm_text(x)
    if not s:
        return ""
    if "sham" in s or "saline" in s:
        return SHAM
    has_medial = ("medial" in s) or ("m+l" in s) or ("m l" in s)
    has_lateral = "lateral" in s
    if (has_medial and has_lateral) or "complete medial" in s or "partial medial" in s:
        return ML
    if "medial and lateral" in s:
        return ML
    if "lateral only" in s or "lateral lesion only" in s or "unilateral" in s:
        return LATERAL
    if has_lateral and not has_medial:
        return LATERAL
    return ""


def choose_group_column(df: pd.DataFrame) -> str:
    candidates = [
        "lesion_group", "lesion_group_detailed", "display_group", "group", "lesion_hit_type"
    ]
    for c in candidates:
        if c in df.columns and df[c].map(normalize_group).ne("").any():
            return c
    raise ValueError(
        "Could not identify a lesion-group column. Tried: "
        + ", ".join(candidates)
        + f"\nFound columns: {list(df.columns)}"
    )


def finite(values: Iterable[object]) -> np.ndarray:
    x = pd.to_numeric(pd.Series(list(values)), errors="coerce").to_numpy(float)
    return x[np.isfinite(x)]


def holm_adjust(p_values: Iterable[float]) -> np.ndarray:
    p = np.asarray(list(p_values), dtype=float)
    out = np.full_like(p, np.nan, dtype=float)
    valid_idx = np.flatnonzero(np.isfinite(p))
    if valid_idx.size == 0:
        return out
    vals = p[valid_idx]
    order = np.argsort(vals)
    m = len(vals)
    adjusted_by_original = np.empty(m, dtype=float)
    running_max = 0.0
    for rank, ord_idx in enumerate(order):
        adj = min(1.0, (m - rank) * vals[ord_idx])
        running_max = max(running_max, adj)
        adjusted_by_original[ord_idx] = running_max
    out[valid_idx] = adjusted_by_original
    return out


def permutation_mean_difference(
    x: np.ndarray,
    y: np.ndarray,
    *,
    alternative: str,
    max_exact_assignments: int,
    monte_carlo_permutations: int,
    rng: np.random.Generator,
) -> dict[str, object]:
    """Test statistic = mean(x) - mean(y)."""
    x = x[np.isfinite(x)]
    y = y[np.isfinite(y)]
    if len(x) < 1 or len(y) < 1:
        return {
            "mean_difference_group1_minus_group2": np.nan,
            "p_value": np.nan,
            "method": "none",
            "assignments_used": 0,
            "possible_exact_assignments": 0,
        }

    observed = float(np.mean(x) - np.mean(y))
    pooled = np.concatenate([x, y])
    n1 = len(x)
    n_total = len(pooled)
    total_assignments = math.comb(n_total, n1)
    eps = 1e-12

    if total_assignments <= max_exact_assignments:
        null = np.empty(total_assignments, dtype=float)
        all_idx = np.arange(n_total)
        for i, combo in enumerate(itertools.combinations(all_idx, n1)):
            mask = np.zeros(n_total, dtype=bool)
            mask[list(combo)] = True
            null[i] = float(np.mean(pooled[mask]) - np.mean(pooled[~mask]))
        if alternative == "greater":
            p = float(np.mean(null >= observed - eps))
        elif alternative == "two-sided":
            p = float(np.mean(np.abs(null) >= abs(observed) - eps))
        else:
            raise ValueError(f"Unknown alternative: {alternative}")
        return {
            "mean_difference_group1_minus_group2": observed,
            "p_value": p,
            "method": "exact_label_permutation",
            "assignments_used": int(total_assignments),
            "possible_exact_assignments": int(total_assignments),
        }

    extreme = 0
    for _ in range(monte_carlo_permutations):
        perm = rng.permutation(pooled)
        stat = float(np.mean(perm[:n1]) - np.mean(perm[n1:]))
        if alternative == "greater":
            extreme += int(stat >= observed - eps)
        elif alternative == "two-sided":
            extreme += int(abs(stat) >= abs(observed) - eps)
    p = float((extreme + 1) / (monte_carlo_permutations + 1))
    return {
        "mean_difference_group1_minus_group2": observed,
        "p_value": p,
        "method": "monte_carlo_label_permutation",
        "assignments_used": int(monte_carlo_permutations),
        "possible_exact_assignments": int(total_assignments),
    }


def load_data(path: Path) -> pd.DataFrame:
    path = path.expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(
            f"Input not found:\n  {path}\nPass the current Figure-3 pair-metrics CSV with --input."
        )
    df = pd.read_csv(path)
    required = {"animal_id", "syllable", "delta_cv"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}\nFound: {list(df.columns)}")
    group_col = choose_group_column(df)
    out = df.copy()
    out["animal_id"] = out["animal_id"].astype(str).str.strip()
    out["delta_cv"] = pd.to_numeric(out["delta_cv"], errors="coerce")
    out["group"] = out[group_col].map(normalize_group)
    out = out.dropna(subset=["delta_cv"])
    out = out[out["group"].isin(GROUP_ORDER)].copy()
    dup = out.duplicated(["animal_id", "syllable"], keep=False)
    if dup.any():
        examples = out.loc[dup, ["animal_id", "syllable"]].head(10)
        raise ValueError(
            "Duplicate animal_id × syllable rows found. Use the Figure-3 balanced pair-metrics table.\n"
            + examples.to_string(index=False)
        )
    return out


def make_bird_quantiles(df: pd.DataFrame, quantile_method: str) -> pd.DataFrame:
    rows = []
    for (bird, group), bdf in df.groupby(["animal_id", "group"], sort=True):
        vals = finite(bdf["delta_cv"])
        for q in QUANTILES:
            qval = float(np.quantile(vals, q, method=quantile_method))
            rows.append({
                "animal_id": bird,
                "group": group,
                "group_label": GROUP_LABEL[group],
                "quantile": q,
                "quantile_label": f"Q{int(round(q * 100))}",
                "quantile_delta_cv": qval,
                "n_qualifying_syllables": int(len(vals)),
            })
    return pd.DataFrame(rows)


def make_group_summary(bird_q: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for q in QUANTILES:
        qdf = bird_q[np.isclose(bird_q["quantile"], q)]
        for group in GROUP_ORDER:
            vals = finite(qdf.loc[qdf["group"] == group, "quantile_delta_cv"])
            rows.append({
                "quantile": q,
                "quantile_label": f"Q{int(round(q * 100))}",
                "group": group,
                "group_label": GROUP_LABEL[group],
                "n_birds": int(len(vals)),
                "mean_bird_quantile_delta_cv": float(np.mean(vals)),
                "median_bird_quantile_delta_cv": float(np.median(vals)),
                "sd_bird_quantile_delta_cv": float(np.std(vals, ddof=1)) if len(vals) > 1 else np.nan,
            })
    return pd.DataFrame(rows)


def make_pairwise_tests(
    bird_q: pd.DataFrame,
    *,
    max_exact_assignments: int,
    monte_carlo_permutations: int,
    seed: int,
) -> pd.DataFrame:
    rows = []
    for qi, q in enumerate(QUANTILES):
        qdf = bird_q[np.isclose(bird_q["quantile"], q)]
        for ci, (comparison, g1, g2, alternative) in enumerate(COMPARISONS):
            x = finite(qdf.loc[qdf["group"] == g1, "quantile_delta_cv"])
            y = finite(qdf.loc[qdf["group"] == g2, "quantile_delta_cv"])
            result = permutation_mean_difference(
                x, y,
                alternative=alternative,
                max_exact_assignments=max_exact_assignments,
                monte_carlo_permutations=monte_carlo_permutations,
                rng=np.random.default_rng(seed + qi * 1000 + ci * 101),
            )
            rows.append({
                "quantile": q,
                "quantile_label": f"Q{int(round(q * 100))}",
                "comparison": comparison,
                "group1": GROUP_LABEL[g1],
                "group2": GROUP_LABEL[g2],
                "alternative": alternative,
                "test_statistic": "difference in mean bird-level quantile ΔCV (group1 - group2)",
                "n_group1": int(len(x)),
                "n_group2": int(len(y)),
                "mean_group1": float(np.mean(x)),
                "mean_group2": float(np.mean(y)),
                "median_group1": float(np.median(x)),
                "median_group2": float(np.median(y)),
                **result,
            })
    tests = pd.DataFrame(rows)

    tests["p_holm_within_quantile"] = np.nan
    for _, idx in tests.groupby("quantile").groups.items():
        idx = list(idx)
        tests.loc[idx, "p_holm_within_quantile"] = holm_adjust(
            tests.loc[idx, "p_value"].to_numpy(float)
        )

    tests["p_holm_across_quantiles_same_comparison"] = np.nan
    for _, idx in tests.groupby("comparison").groups.items():
        idx = list(idx)
        tests.loc[idx, "p_holm_across_quantiles_same_comparison"] = holm_adjust(
            tests.loc[idx, "p_value"].to_numpy(float)
        )

    tests["p_holm_all_9_sensitivity_tests"] = holm_adjust(
        tests["p_value"].to_numpy(float)
    )
    return tests


def make_compact_table(group_summary: pd.DataFrame, tests: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for q in QUANTILES:
        qlab = f"Q{int(round(q * 100))}"
        g = group_summary[np.isclose(group_summary["quantile"], q)]
        t = tests[np.isclose(tests["quantile"], q)]

        def gval(group: str, col: str):
            r = g[g["group"] == group]
            return r.iloc[0][col] if not r.empty else np.nan

        def tval(comp: str, col: str):
            r = t[t["comparison"] == comp]
            return r.iloc[0][col] if not r.empty else np.nan

        rows.append({
            "Quantile": qlab,
            "Sham n": int(gval(SHAM, "n_birds")),
            "Sham median ΔCV": gval(SHAM, "median_bird_quantile_delta_cv"),
            "Lateral-only n": int(gval(LATERAL, "n_birds")),
            "Lateral-only median ΔCV": gval(LATERAL, "median_bird_quantile_delta_cv"),
            "Medial+lateral n": int(gval(ML, "n_birds")),
            "Medial+lateral median ΔCV": gval(ML, "median_bird_quantile_delta_cv"),
            "M+L - Sham mean difference": tval("M+L vs Sham", "mean_difference_group1_minus_group2"),
            "M+L vs Sham exact p": tval("M+L vs Sham", "p_value"),
            "M+L vs Sham Holm p (within Q)": tval("M+L vs Sham", "p_holm_within_quantile"),
            "M+L vs Sham Holm p (across Q60/Q75/Q90)": tval(
                "M+L vs Sham", "p_holm_across_quantiles_same_comparison"
            ),
            "M+L - Lateral mean difference": tval(
                "M+L vs Lateral-only", "mean_difference_group1_minus_group2"
            ),
            "M+L vs Lateral exact p": tval("M+L vs Lateral-only", "p_value"),
            "M+L vs Lateral Holm p (within Q)": tval(
                "M+L vs Lateral-only", "p_holm_within_quantile"
            ),
            "Lateral - Sham mean difference": tval(
                "Lateral-only vs Sham", "mean_difference_group1_minus_group2"
            ),
            "Lateral vs Sham exact two-sided p": tval("Lateral-only vs Sham", "p_value"),
            "Lateral vs Sham Holm p (within Q)": tval(
                "Lateral-only vs Sham", "p_holm_within_quantile"
            ),
        })
    return pd.DataFrame(rows)



def fmt_value(x: object, digits: int = 3) -> str:
    """Compact manuscript-facing numeric formatting."""
    try:
        value = float(x)
    except (TypeError, ValueError):
        return ""
    if not np.isfinite(value):
        return ""
    return f"{value:.{digits}f}"


def fmt_p(x: object) -> str:
    """Format p-values without implying more precision than useful."""
    try:
        p = float(x)
    except (TypeError, ValueError):
        return ""
    if not np.isfinite(p):
        return ""
    if p < 0.0001:
        return "<0.0001"
    return f"{p:.4f}"


def _set_cell_text(cell, text: str, *, bold: bool = False, font_size: float = 9.0) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def make_word_supplement(
    *,
    compact: pd.DataFrame,
    tests: pd.DataFrame,
    group_summary: pd.DataFrame,
    output_path: Path,
    input_path: Path,
) -> None:
    """Write a manuscript-ready DOCX containing the Q60/Q75/Q90 sensitivity results."""
    doc = Document()

    # Landscape makes the main sensitivity table easy to copy/paste without cramped columns.
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run(
        "Supplemental Table X. Sensitivity of the phrase-duration variability result "
        "to the bird-level quantile summary."
    )
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    rr = intro.add_run(
        "For each bird, the 60th, 75th, or 90th percentile (Q60, Q75, Q90) of syllable-level "
        "ΔCV (post-lesion minus late pre-lesion phrase-duration CV) was calculated across all "
        "Figure 3-qualifying syllables. Values below are bird-level group medians; p values are "
        "Holm-adjusted within each quantile across the three planned group comparisons."
    )
    rr.font.name = "Arial"
    rr.font.size = Pt(9.5)

    headers = [
        "Quantile",
        "Sham\nmedian ΔCV",
        "Lateral-only\nmedian ΔCV",
        "Medial+lateral\nmedian ΔCV",
        "M+L vs Sham\nHolm-adjusted p",
        "M+L vs Lateral-only\nHolm-adjusted p",
        "Lateral-only vs Sham\nHolm-adjusted p",
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for j, h in enumerate(headers):
        _set_cell_text(table.rows[0].cells[j], h, bold=True, font_size=8.5)

    for _, row in compact.iterrows():
        cells = table.add_row().cells
        values = [
            row["Quantile"],
            fmt_value(row["Sham median ΔCV"]),
            fmt_value(row["Lateral-only median ΔCV"]),
            fmt_value(row["Medial+lateral median ΔCV"]),
            fmt_p(row["M+L vs Sham Holm p (within Q)"]),
            fmt_p(row["M+L vs Lateral Holm p (within Q)"]),
            fmt_p(row["Lateral vs Sham Holm p (within Q)"]),
        ]
        for j, value in enumerate(values):
            _set_cell_text(cells[j], value, font_size=9.0)

    # Derive n directly rather than hard-coding manuscript sample sizes.
    ns = {}
    for group in GROUP_ORDER:
        rows = group_summary[group_summary["group"] == group]
        if not rows.empty:
            ns[group] = int(rows.iloc[0]["n_birds"])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run(
        "Note. Bird is the independent biological replicate. The permutation-test statistic was "
        "the difference in mean bird-level quantile ΔCV between groups. M+L vs Sham and M+L vs "
        "Lateral-only were one-sided tests with the alternative M+L > comparison group; "
        "Lateral-only vs Sham was two-sided. Exact bird-label permutations were used whenever "
        "feasible. Sample sizes: "
        f"Sham n={ns.get(SHAM, 'NA')}, Lateral-only n={ns.get(LATERAL, 'NA')}, "
        f"Medial+lateral n={ns.get(ML, 'NA')}."
    )
    note_run.font.name = "Arial"
    note_run.font.size = Pt(8.5)

    # Add a compact detailed table for audit/reviewer response; the first table is the
    # one intended for direct manuscript copy/paste.
    hdr = doc.add_paragraph()
    hdr.paragraph_format.space_before = Pt(3)
    hdr.paragraph_format.space_after = Pt(4)
    hr = hdr.add_run("Detailed exact permutation results")
    hr.bold = True
    hr.font.name = "Arial"
    hr.font.size = Pt(10)

    detailed_headers = [
        "Quantile",
        "Comparison",
        "Alternative",
        "Mean difference",
        "Raw p",
        "Holm p\nwithin quantile",
        "Holm p\nacross Q60/Q75/Q90",
    ]
    dtable = doc.add_table(rows=1, cols=len(detailed_headers))
    dtable.style = "Table Grid"
    dtable.alignment = WD_TABLE_ALIGNMENT.CENTER
    dtable.autofit = True
    for j, h in enumerate(detailed_headers):
        _set_cell_text(dtable.rows[0].cells[j], h, bold=True, font_size=8.0)

    for _, row in tests.sort_values(["quantile", "comparison"]).iterrows():
        cells = dtable.add_row().cells
        alt = ">" if row["alternative"] == "greater" else "two-sided"
        values = [
            row["quantile_label"],
            row["comparison"],
            alt,
            fmt_value(row["mean_difference_group1_minus_group2"], digits=3),
            fmt_p(row["p_value"]),
            fmt_p(row["p_holm_within_quantile"]),
            fmt_p(row["p_holm_across_quantiles_same_comparison"]),
        ]
        for j, value in enumerate(values):
            _set_cell_text(cells[j], value, font_size=8.0)

    provenance = doc.add_paragraph()
    provenance.paragraph_format.space_before = Pt(5)
    pr = provenance.add_run(f"Analysis input: {input_path.resolve()}")
    pr.italic = True
    pr.font.name = "Arial"
    pr.font.size = Pt(7.5)

    doc.save(output_path)

def main() -> None:
    args = parse_args()
    args.input = args.input.expanduser()
    args.output = args.output.expanduser()
    args.output.mkdir(parents=True, exist_ok=True)

    df = load_data(args.input)
    bird_q = make_bird_quantiles(df, args.quantile_method)
    group_summary = make_group_summary(bird_q)
    tests = make_pairwise_tests(
        bird_q,
        max_exact_assignments=args.max_exact_assignments,
        monte_carlo_permutations=args.monte_carlo_permutations,
        seed=args.seed,
    )
    compact = make_compact_table(group_summary, tests)

    bird_q.to_csv(args.output / "01_bird_level_Q60_Q75_Q90.csv", index=False)
    group_summary.to_csv(args.output / "02_group_summary_Q60_Q75_Q90.csv", index=False)
    tests.to_csv(args.output / "03_pairwise_exact_permutation_tests.csv", index=False)
    compact.to_csv(args.output / "04_supplemental_sensitivity_table.csv", index=False)

    word_path = args.output / "05_Supplemental_Table_Q60_Q75_Q90_sensitivity.docx"
    make_word_supplement(
        compact=compact,
        tests=tests,
        group_summary=group_summary,
        output_path=word_path,
        input_path=args.input,
    )

    readme = [
        "Q60 / Q75 / Q90 phrase-duration ΔCV sensitivity analysis",
        "========================================================",
        "",
        f"Input: {args.input.resolve()}",
        f"Qualifying bird × syllable rows: {len(df)}",
        f"Birds: {df['animal_id'].nunique()}",
        "",
        "Per-bird summary: Q60, Q75, or Q90 of syllable-level ΔCV.",
        "Independent biological replicate: bird.",
        "Permutation statistic: difference in mean bird-level quantile ΔCV.",
        "M+L > Sham and M+L > lateral-only are one-sided.",
        "Lateral-only vs Sham is two-sided.",
        "Holm-within-quantile matches the manuscript-style 3-comparison family.",
        "Additional across-quantile Holm columns are conservative sensitivity checks.",
        "",
        "M+L vs Sham results:",
    ]
    sub = tests[tests["comparison"] == "M+L vs Sham"].sort_values("quantile")
    for _, r in sub.iterrows():
        readme.append(
            f"  {r['quantile_label']}: mean diff={r['mean_difference_group1_minus_group2']:.6g}; "
            f"p={r['p_value']:.6g}; Holm-within-Q={r['p_holm_within_quantile']:.6g}; "
            f"Holm-across-Q={r['p_holm_across_quantiles_same_comparison']:.6g}"
        )
    (args.output / "06_readme.txt").write_text("\n".join(readme) + "\n", encoding="utf-8")

    print("Input:")
    print(f"  {args.input.resolve()}")
    print(f"Qualifying bird × syllable rows: {len(df)}")
    print(f"Birds: {df['animal_id'].nunique()}")
    print()
    print("M+L vs Sham sensitivity:")
    print(
        sub[[
            "quantile_label",
            "mean_difference_group1_minus_group2",
            "p_value",
            "p_holm_within_quantile",
            "p_holm_across_quantiles_same_comparison",
        ]].to_string(index=False)
    )
    print()
    print("Saved:")
    for name in [
        "01_bird_level_Q60_Q75_Q90.csv",
        "02_group_summary_Q60_Q75_Q90.csv",
        "03_pairwise_exact_permutation_tests.csv",
        "04_supplemental_sensitivity_table.csv",
        "05_Supplemental_Table_Q60_Q75_Q90_sensitivity.docx",
        "06_readme.txt",
    ]:
        print(f"  {args.output / name}")


if __name__ == "__main__":
    main()
