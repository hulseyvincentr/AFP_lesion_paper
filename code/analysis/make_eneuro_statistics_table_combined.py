#!/usr/bin/env python3
"""Generate manuscript-matched tests, bird-level percentile bootstrap CIs and Word tables.
Dependencies: numpy, pandas, scipy, python-docx. No annotation pipeline rerun needed.
Intervals are pointwise 95% percentile intervals, not multiplicity-adjusted or
inversions of the hypothesis tests. Small bird samples limit interval accuracy.
"""
import argparse, itertools, math
from pathlib import Path
import numpy as np
import pandas as pd
from scipy.stats import wilcoxon
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def holm(p):
 p=np.asarray(p,float); order=np.argsort(p); out=np.empty(len(p));out[order]=np.minimum(1,np.maximum.accumulate(p[order]*(len(p)-np.arange(len(p)))));return out

def permutation(x,y,two_sided):
 observed=x.mean()-y.mean(); pooled=np.r_[x,y]; count=0
 for indices in itertools.combinations(range(len(pooled)),len(x)):
  mask=np.zeros(len(pooled),bool);mask[list(indices)]=True
  value=pooled[mask].mean()-pooled[~mask].mean()
  count+=int(abs(value)>=abs(observed)-1e-12 if two_sided else value>=observed-1e-12)
 return count/math.comb(len(pooled),len(x))

def interval(x,y,rng,reps,paired=False):
 # Paired effects bootstrap entire bird differences, never separate pre/post samples.
 draws=[]
 for start in range(0,reps,10000):
  size=min(10000,reps-start)
  if paired:
   samples=rng.choice(x,size=(size,len(x)),replace=True);draws.append(np.median(samples,axis=1))
  else:
   a=rng.choice(x,size=(size,len(x)),replace=True).mean(axis=1)
   b=rng.choice(y,size=(size,len(y)),replace=True).mean(axis=1);draws.append(a-b)
 return np.quantile(np.concatenate(draws),[.025,.975])

def checked(df,keys,columns):
 if df.duplicated(keys).any():raise ValueError('Duplicate bird records: '+str(keys))
 if not np.isfinite(df[columns].to_numpy(float)).all():raise ValueError('Non-finite input values')

def main():
 p=argparse.ArgumentParser(description=__doc__)
 p.add_argument('--quantiles',type=Path,required=True);p.add_argument('--permutation-results',type=Path,required=True)
 p.add_argument('--panel-c',type=Path,required=True);p.add_argument('--panel-d',type=Path,required=True)
 p.add_argument('--outdir',type=Path,required=True);p.add_argument('--seed',type=int,default=123)
 p.add_argument('--bootstrap-reps',type=int,default=200000)
 a=p.parse_args()
 if a.bootstrap_reps<10000:raise ValueError('Use at least 10000 bootstrap draws')
 q=pd.read_csv(a.quantiles);tests=pd.read_csv(a.permutation_results);c=pd.read_csv(a.panel_c);d=pd.read_csv(a.panel_d)
 checked(q,['animal_id','quantile'],['quantile_delta_cv']);checked(c,['animal_id'],['pre_bc','post_bc']);checked(d,['animal_id'],['q75_pre_bc','q75_post_bc','below_q75_pre_bc','below_q75_post_bc'])
 rng=np.random.default_rng(a.seed); rows=[]
 for quantile in [.75,.60,.90]:
  group=q[np.isclose(q['quantile'],quantile)]
  family=[]
  for g1,g2,two in [('Medial+lateral','Sham',False),('Medial+lateral','Lateral-only',False),('Lateral-only','Sham',True)]:
   x=group.loc[group.group_label==g1,'quantile_delta_cv'].to_numpy(float);y=group.loc[group.group_label==g2,'quantile_delta_cv'].to_numpy(float)
   expected={'Medial+lateral':9,'Lateral-only':8,'Sham':4}
   if len(x)!=expected[g1] or len(y)!=expected[g2]:raise ValueError('Unexpected group size')
   source=tests[np.isclose(tests['quantile'],quantile)&(tests.group1==g1)&(tests.group2==g2)]
   if len(source)!=1:raise ValueError('Missing or duplicate reference comparison')
   r=source.iloc[0];effect=x.mean()-y.mean();raw=permutation(x,y,two)
   if not np.isclose(effect,r.mean_difference_group1_minus_group2,atol=1e-10) or not np.isclose(raw,r.p_value,atol=1e-10):raise ValueError('Quantile inputs do not reproduce reference results')
   low,high=interval(x,y,rng,a.bootstrap_reps)
   family.append(dict(figure='3B' if quantile==.75 else 'Extended Data Table 3-2',comparison=f'Q{round(quantile*100)} {g1} minus {g2}',data_structure='Independent bird-level quantile ΔCV',n=f'{len(x)} vs {len(y)}',test='Exact permutation, '+('two-sided' if two else 'one-sided'),effect_measure='Difference in means',effect=effect,ci_low=low,ci_high=high,test_statistic=effect,p_raw=raw,reference_adjusted=r.p_holm_within_quantile,correction='Holm across 3 comparisons within quantile'))
  for r,adjusted in zip(family,holm([r['p_raw'] for r in family])):
   if not np.isclose(adjusted,r['reference_adjusted'],atol=1e-10):raise ValueError('Holm correction mismatch')
   r['p_adjusted']=adjusted;rows.append(r)
 bc=[]
 for name,g in c.groupby('pooled_group',sort=False):
  label='Medial+lateral' if 'medial' in name.lower() else ('Sham' if 'sham' in name.lower() else 'Lateral-only')
  bc.append(('4C',label,g.pre_bc.to_numpy(float),g.post_bc.to_numpy(float)))
 for prefix,label in [('q75','Medial+lateral ≥Q75'),('below_q75','Medial+lateral <Q75')]:bc.append(('4D',label,d[prefix+'_pre_bc'].to_numpy(float),d[prefix+'_post_bc'].to_numpy(float)))
 bcrows=[]
 for panel,label,pre,post in bc:
  diff=post-pre;result=wilcoxon(post,pre,alternative='two-sided',zero_method='wilcox',method='auto');low,high=interval(diff,None,rng,a.bootstrap_reps,paired=True)
  bcrows.append(dict(figure=panel,comparison=label+' post minus pre',data_structure='Paired bird-level median BC',n=str(len(diff)),test='Wilcoxon signed-rank, two-sided',effect_measure='Median paired change',effect=np.median(diff),ci_low=low,ci_high=high,test_statistic=result.statistic,p_raw=result.pvalue,correction='Holm within panel'))
 for panel in ['4C','4D']:
  family=[r for r in bcrows if r['figure']==panel]
  for r,adj in zip(family,holm([r['p_raw'] for r in family])):r['p_adjusted']=adj
 # Reorder only after calculation so seed 123 retains the original CIs.
 # Results order: Q75, sensitivity Q60/Q90, Figure 4C M+L/sham/lateral,
 # then Figure 4D >=Q75/<Q75. Letters are assigned after this ordering.
 panel_c_order={'Medial+lateral post minus pre':0,'Sham post minus pre':1,'Lateral-only post minus pre':2}
 ordered_c=sorted([r for r in bcrows if r['figure']=='4C'],key=lambda r:panel_c_order[r['comparison']])
 rows=rows+ordered_c+[r for r in bcrows if r['figure']=='4D']
 for i,r in enumerate(rows):r['row']=chr(97+i)
 a.outdir.mkdir(parents=True,exist_ok=True)
 pd.DataFrame(rows).to_csv(a.outdir/'statistics_full_precision.csv',index=False)
 pd.DataFrame(rows)[['row','figure','comparison','p_raw','p_adjusted']].to_csv(a.outdir/'statistics_row_references.csv',index=False)
 doc=Document();s=doc.sections[0];s.orientation=1;s.page_width=Inches(11);s.page_height=Inches(8.5);s.top_margin=s.bottom_margin=Inches(.6);s.left_margin=s.right_margin=Inches(.6)
 doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.line_spacing=2
 doc.add_paragraph('Table 1. Statistical analyses')
 t=doc.add_table(rows=1,cols=7);t.autofit=False
 # A row-letter stub precedes the three substantive columns requested by eNeuro.
 headings=['','Data structure','Statistical test','95% confidence interval','Comparison, figure and n','Effect estimate','Holm-adjusted p']
 widths=[.35,1.4,1.85,1.4,2.9,1.0,.9]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for cell,h,w in zip(t.rows[0].cells,headings,widths):cell.text=h;cell.width=Inches(w)
 header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
 for r in rows:
  vals=[r['row'],r['data_structure'],r['test'],f"[{r['ci_low']:+.4f}, {r['ci_high']:+.4f}]",r['comparison']+'; '+r['figure']+'; n='+r['n'],f"{r['effect']:+.4f}",f"{r['p_adjusted']:.4f}"]
  for cell,v,w in zip(t.add_row().cells,vals,widths):cell.text=v;cell.width=Inches(w)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();e=OxmlElement('w:cantSplit');pr.append(e)
  for cell in row.cells:
   # Explicit formatting survives copying into a manuscript with other defaults.
   for paragraph in cell.paragraphs:
    paragraph.paragraph_format.line_spacing=2
    paragraph.paragraph_format.space_before=Pt(0)
    paragraph.paragraph_format.space_after=Pt(0)
    for run in paragraph.runs:
     run.font.name='Arial';run.font.size=Pt(10);run.bold=False;run.italic=False
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'AAAAAA');borders.append(e)
   cell._tc.get_or_add_tcPr().append(borders)
 # Keep the legend separate for placement after manuscript references.
 doc.add_page_break()
 doc.add_paragraph('Table 1. Statistical analyses')
 doc.add_paragraph('Lowercase row letters identify the corresponding statistical results in the manuscript. Birds are the independent biological replicates. CV, coefficient of variation; BC, Bhattacharyya coefficient; Q60, Q75 and Q90, the 60th, 75th and 90th percentiles of each bird’s syllable-level ΔCV distribution. ΔCV and ΔBC are post-lesion minus pre-lesion changes. Between-group effects are differences in mean bird-level quantile ΔCV; acoustic effects are medians of bird-level paired ΔBC.')
 doc.add_paragraph(f'Confidence intervals are pointwise 95% percentile bootstrap intervals from {a.bootstrap_reps:,} resamples with replacement (seed {a.seed}). Between-group intervals resample birds independently within each group, preserving group sizes. Paired intervals resample bird-level differences and estimate their median. Holm correction applies to p-values only, across three group contrasts separately within each quantile, three tests in Figure 4C, and two tests in Figure 4D. Intervals are not adjusted for multiple comparisons. Median-change intervals summarize effects separately from the Wilcoxon signed-rank tests. Small samples, especially sham n=4, limit bootstrap precision.')
 doc.save(a.outdir/'eNeuro_statistics_table.docx')
 print(pd.DataFrame(rows)[['row','figure','effect','ci_low','ci_high','p_adjusted']].to_string(index=False))
if __name__=='__main__':main()
